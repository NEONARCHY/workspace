import io

# Russian fixture text intentionally tests Windows-1251.
# ruff: noqa: RUF001
import struct
from zipfile import ZipFile

import openpyxl  # type: ignore[import-untyped]
import pytest

from yuksalish_api.survey_documents import MAX_FILE_BYTES, SurveyDocumentError, parse_workbook
from yuksalish_api.survey_schemas import (
    SurveyConfiguration,
    SurveyFilter,
    SurveySheet,
    SurveyWorkbook,
)
from yuksalish_api.survey_service import validate_configuration


def xlsx_bytes() -> bytes:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Answers"
    sheet.append(["Group", "Amount"])
    sheet.append(["A", 20])
    sheet.append(["B", "=SUM(B2:B2)"])
    sheet.append([None, None])
    sheet.append(["C", 0])
    book.create_sheet("Other").append(["Other survey"])
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def test_xlsx_preserves_sheets_coordinates_and_never_evaluates_formulas() -> None:
    workbook = parse_workbook("answers.xlsx", xlsx_bytes())
    assert [sheet.name for sheet in workbook.sheets] == ["Answers", "Other"]
    assert workbook.sheets[0].rows[2][1] == "=SUM(B2:B2)"
    assert not any(workbook.sheets[0].rows[3])
    assert workbook.sheets[0].rows[4][1] == "0"
    assert "(1)" in workbook.warnings[0]


def test_csv_explicit_encoding_delimiter_quotes_and_zero() -> None:
    content = 'Вопрос;Число\n"a;b";0\nнет;null\n'.encode("cp1251")
    result = parse_workbook("answers.csv", content, encoding="cp1251", delimiter=";")
    assert result.sheets[0].rows == [["Вопрос", "Число"], ["a;b", "0"], ["нет", "null"]]
    with pytest.raises(SurveyDocumentError):
        parse_workbook("answers.csv", content)


@pytest.mark.parametrize("filename, content", [
    ("answers.xlsm", b"data"), ("answers.xls", b"not-an-xls"),
    ("answers.xlsx", b"not-a-zip"), ("answers.csv", b""),
    ("answers.csv", b" \n \n"), ("answers.csv", b"x" * (MAX_FILE_BYTES + 1)),
], ids=["macro-extension", "invalid-xls", "invalid-xlsx", "empty", "blanks", "oversized"])
def test_invalid_inputs_are_safe_validation_errors(filename: str, content: bytes) -> None:
    with pytest.raises(SurveyDocumentError):
        parse_workbook(filename, content)


@pytest.mark.parametrize("content", [
    ("a;" * 513).encode(), ("row\n" * 25_001).encode(), b"x" * 140_000,
], ids=["too-wide", "too-tall", "oversized-cell"])
def test_csv_limits_reject_without_silent_truncation(content: bytes) -> None:
    with pytest.raises(SurveyDocumentError):
        parse_workbook("data.csv", content, delimiter=";")


def test_sparse_matrix_cannot_expand_to_an_unbounded_export() -> None:
    # Few actual values, but too many reconstructed cells for browser/export safety.
    content = (";".join(["Q"] * 500) + "\n" + "A\n" * 400).encode()
    with pytest.raises(SurveyDocumentError, match="150 000"):
        parse_workbook("sparse.csv", content, delimiter=";")


@pytest.mark.parametrize("replacement", [
    b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
    b'<sheetData><row r="25001"><c r="A25001"><v>1</v></c></row></sheetData></worksheet>',
    b'<!DOCTYPE x [<!ENTITY foo SYSTEM "file:///secret">]><worksheet>&foo;</worksheet>',
    b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
    b'<sheetData><row r="1"><c r="A1"><v>1</v></c><c r="A1"><v>2</v></c>'
    b'</row></sheetData></worksheet>',
])
def test_sparse_bombs_and_unsafe_xml_are_rejected_before_openpyxl(replacement: bytes) -> None:
    output = io.BytesIO()
    with ZipFile(io.BytesIO(xlsx_bytes())) as source, ZipFile(output, "w") as target:
        for entry in source.infolist():
            content = (replacement if entry.filename == "xl/worksheets/sheet1.xml"
                       else source.read(entry.filename))
            target.writestr(entry.filename, content)
    with pytest.raises(SurveyDocumentError):
        parse_workbook("data.xlsx", output.getvalue())


def test_macro_and_external_link_archives_are_rejected() -> None:
    for filename in ["xl/vbaProject.bin", "xl/externalLinks/externalLink1.xml"]:
        output = io.BytesIO()
        with ZipFile(io.BytesIO(xlsx_bytes())) as source, ZipFile(output, "w") as target:
            for entry in source.infolist():
                target.writestr(entry.filename, source.read(entry.filename))
            target.writestr(filename, b"ignored")
        with pytest.raises(SurveyDocumentError):
            parse_workbook("data.xlsx", output.getvalue())


def test_actual_legacy_xls_biff2_stream() -> None:
    # Synthetic BIFF2, not a real response file. xlrd supports unwrapped BIFF streams.
    def record(code: int, payload: bytes) -> bytes:
        return struct.pack("<HH", code, len(payload)) + payload

    content = record(0x09, struct.pack("<HH", 0x02, 0x10))
    content += record(0x04, struct.pack("<HH3sB", 0, 0, b"\x00\x00\x00", 1) + b"Q")
    content += record(0x03, struct.pack("<HH3sd", 1, 0, b"\x00\x00\x00", 20.0))
    content += record(0x0A, b"")
    result = parse_workbook("sample.xls", content)
    assert result.sheets[0].rows == [["Q"], ["20"]]


def test_excel_export_has_literal_cells_and_validated_source_rows() -> None:
    from yuksalish_api.errors import WorkspaceRepositoryError
    from yuksalish_api.survey_export import export_xlsx
    from yuksalish_api.survey_schemas import SurveyExport

    source = SurveyWorkbook(sheets=[SurveySheet(name="Data", rows=[
        ["=Header"], ["=HYPERLINK(1)"], ["@SUM(1)"],
    ])])
    config = SurveyConfiguration(query="test")
    content = export_xlsx(source, SurveyExport(configuration=config, row_numbers=[2, 3]))
    book = openpyxl.load_workbook(io.BytesIO(content), data_only=False)
    try:
        assert book["Answers"]["B1"].value == "=Header"
        assert book["Answers"]["B2"].value == "=HYPERLINK(1)"
        assert book["Answers"]["B2"].data_type == "s"
        assert book["Methodology"]["B1"].value == "Data"
    finally:
        book.close()
    for selected in [[0], [1], [4], [2, 2]]:
        with pytest.raises(WorkspaceRepositoryError):
            export_xlsx(source, SurveyExport(configuration=config, row_numbers=selected))


@pytest.mark.parametrize("configuration", [
    SurveyConfiguration(sheet=1), SurveyConfiguration(header_row=2),
    SurveyConfiguration(types={"100": "numeric"}),
    SurveyConfiguration(types={"not-a-column": "text"}),
    SurveyConfiguration(types={"9" * 5000: "text"}),
    SurveyConfiguration(filters=[SurveyFilter(column=2, value="x")]),
    SurveyConfiguration(missing_values=["x" * 201]),
])
def test_invalid_settings_cannot_reference_nonexistent_source(
    configuration: SurveyConfiguration,
) -> None:
    from yuksalish_api.errors import WorkspaceRepositoryError

    workbook = SurveyWorkbook(sheets=[SurveySheet(name="data", rows=[["Q"], ["A"]])])
    with pytest.raises(WorkspaceRepositoryError) as caught:
        validate_configuration(configuration, workbook)
    assert caught.value.status_code == 422
