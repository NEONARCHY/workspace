import asyncio
from datetime import UTC, datetime
from io import BytesIO
from unittest.mock import AsyncMock
from uuid import uuid4
from zipfile import ZipFile

import openpyxl
import pytest
from defusedxml.ElementTree import fromstring
from fastapi import FastAPI, HTTPException, Response
from httpx import ASGITransport, AsyncClient
from PIL import Image
from pydantic import ValidationError
from sqlalchemy import select, update

from test_personnel_dossiers import data, scenario
from yuksalish_api import dossier_service, photo_roster_service
from yuksalish_api.auth import require_user
from yuksalish_api.database import get_connection
from yuksalish_api.dossier_documents import portrait_jpeg
from yuksalish_api.dossier_schemas import (
    DossierAccessRule,
    DossierAccessWrite,
    PhotoCard,
    PhotoRosterExport,
)
from yuksalish_api.errors import WorkspaceRepositoryError
from yuksalish_api.object_storage import ObjectStorageError
from yuksalish_api.photo_roster_documents import W, render_excel, render_word
from yuksalish_api.routers import dossiers as routes
from yuksalish_api.tables import audit_events, hr_settings, users


def portrait():
    output = BytesIO()
    Image.new("RGB", (120, 160), "#293A55").save(output, "JPEG")
    return output.getvalue()


def test_roster_contract_is_bounded_and_cannot_supply_personal_values():
    identity = uuid4()
    for payload in ({"userIds": []}, {"userIds": [identity, identity]},
                    {"userIds": [uuid4() for _ in range(501)]},
                    {"userIds": [identity], "fullName": "Override"}):
        with pytest.raises(ValidationError):
            PhotoRosterExport.model_validate(payload)


def test_documents_have_identical_minimal_fields_photos_and_literal_excel_values():
    cards = [PhotoCard(user_id=uuid4(), full_name="=HYPERLINK(\"https://example.invalid\")",
                       job_title="Бош мутахассис <&>\nЯн\u0433и қатор", department_id=None,
                       birth_year=1991, universities=["Тестовый университет", "Ҳуқуқ академияси"]),
             PhotoCard(user_id=uuid4(), full_name="Сотрудник без анкеты", job_title=None,
                       department_id=None, birth_year=None, universities=[])]
    photos = [portrait_jpeg(portrait()), None]
    now = datetime(2026, 10, 22, 18, 0, tzinfo=UTC)
    word = render_word(cards, photos, now)
    with ZipFile(BytesIO(word)) as archive:
        root = fromstring(archive.read("word/document.xml"))
        text = " ".join(root.itertext())
        assert all(value in text for value in ("1991", "Ҳуқуқ академияси", "Бош мутахассис <&>",
                                               "Сотрудник без анкеты", "Нет фото"))
        assert len(root.findall(f".//{{{W}}}tblHeader")) == 1
        assert root.find(f".//{{{W}}}pgSz").attrib[f"{{{W}}}orient"] == "landscape"
        assert len([name for name in archive.namelist() if name.startswith("word/media/")]) == 1
        assert b"External" not in archive.read("word/_rels/document.xml.rels")
        default_namespace = b'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"'
        assert default_namespace in archive.read("[Content_Types].xml")
        assert b"Title" in archive.read("word/styles.xml")
    book = openpyxl.load_workbook(BytesIO(render_excel(cards, photos, now)), data_only=False)
    sheet = book.active
    assert sheet["A5"].value == 1 and sheet["D5"].value == 1991
    assert sheet["C5"].data_type == "s" and sheet["C5"].value == cards[0].full_name
    assert sheet["E5"].value == "Тестовый университет\nҲ\u0443қ\u0443қ академияси"
    assert sheet["B6"].value == "Нет фото" and sheet["D6"].value != 0
    assert sheet.freeze_panes == "D5" and sheet.auto_filter.ref == "A4:F6"
    assert sheet.print_title_rows == "$1:$4" and len(sheet._images) == 1
    assert not any(cell.data_type == "f" for row in sheet for cell in row)
    book.close()


def test_long_education_continues_with_employee_identity_and_without_clipped_excel_rows():
    institutions = [f"Institution {n:02d} " + "International university " * 25
                    for n in range(20)]
    card = PhotoCard(user_id=uuid4(), full_name="Long education employee", job_title="Specialist",
                     department_id=None, birth_year=1990, universities=institutions)
    now = datetime.now(UTC)
    word = render_word([card], [None], now)
    with ZipFile(BytesIO(word)) as archive:
        root = fromstring(archive.read("word/document.xml"))
        text = " ".join(root.itertext())
        assert "Institution 00" in text and "Institution 19" in text
        assert text.count(card.full_name) > 1
    book = openpyxl.load_workbook(BytesIO(render_excel([card], [None], now)))
    sheet = book.active
    text = "\n".join(str(row[4].value or "") for row in list(sheet)[4:])
    assert "Institution 00" in text and "Institution 19" in text
    assert sheet.max_row > 5
    assert all(sheet.cell(row, 1).value == 1 for row in range(5, sheet.max_row + 1))
    assert all(sheet.row_dimensions[row].height < 409 for row in range(5, sheet.max_row + 1))
    book.close()


def test_exports_are_available_to_hr_deputy_chair_and_admin_but_not_employee(monkeypatch):
    async def run():
        async with scenario() as (connection, actors):
            employee = actors["employee"]
            await dossier_service.save(connection, employee, employee.id, data())
            await connection.execute(update(users).where(users.c.id == employee.id).values(
                avatar_storage_key="qa-private-photo",
            ))
            app = FastAPI()
            app.include_router(routes.router, prefix="/api/v1")
            app.state.object_storage = AsyncMock()
            app.state.object_storage.get.return_value = portrait()
            app.state.settings = object()

            async def test_connection():
                yield connection

            app.dependency_overrides[get_connection] = test_connection
            actor = employee
            app.dependency_overrides[require_user] = lambda: actor
            convert = AsyncMock(return_value=Response(b"%PDF-qa", media_type="application/pdf"))
            monkeypatch.setattr(routes, "render_word_pdf", convert)
            async with AsyncClient(
                transport=ASGITransport(app=app), base_url="http://test",
            ) as client:
                payload = {"userIds": [str(employee.id)]}
                for kind in ("word", "excel", "pdf"):
                    denied = await client.post(
                        f"/api/v1/personnel-dossiers/photos/export/{kind}", json=payload,
                    )
                    assert denied.status_code == 403
                app.state.object_storage.get.assert_not_called()
                for role in ("hr", "deputy", "admin", "employee"):
                    actor = actors[role]
                    if role == "employee":
                        await connection.execute(update(hr_settings).values(chair_user_id=actor.id))
                    for kind, extension in (("word", "docx"), ("excel", "xlsx"), ("pdf", "pdf")):
                        result = await client.post(
                            f"/api/v1/personnel-dossiers/photos/export/{kind}", json=payload,
                        )
                        assert result.status_code == 200, result.text[:200]
                        assert result.headers["cache-control"] == "private, no-store"
                        assert result.headers["x-content-type-options"] == "nosniff"
                        assert f"photo-roster.{extension}" in result.headers["content-disposition"]
                        if kind == "word":
                            with ZipFile(BytesIO(result.content)) as archive:
                                xml = archive.read("word/document.xml")
                                assert b"PRIVATE BIRTHPLACE" not in xml
                                assert b"PRIVATE NOTES" not in xml
                                assert b"1991-05-12" not in xml
                assert convert.await_count == 4
                events = (await connection.execute(select(audit_events.c.details).where(
                    audit_events.c.action == "dossier.photo_roster_exported",
                ))).scalars().all()
                assert len(events) == 12
                assert all(set(event) == {"format", "count"} for event in events)
    asyncio.run(run())


def test_scoped_access_no_silent_partial_export_and_revocation_during_conversion(monkeypatch):
    async def run():
        async with scenario() as (connection, actors):
            employee, other, admin = actors["employee"], actors["deputy"], actors["admin"]
            await dossier_service.save_access(connection, admin, DossierAccessWrite(
                revision=0, reason="QA scope", rules=[DossierAccessRule(
                    grantee_user_id=employee.id, scope="employee", scope_id=other.id,
                    photo_roster=True,
                )],
            ))
            storage = AsyncMock()
            scoped = await photo_roster_service.source(connection, employee,
                PhotoRosterExport(user_ids=[other.id]), storage)
            assert [card.user_id for card in scoped.cards] == [other.id]
            storage.get.assert_not_called()
            for identities in ([other.id, admin.id], [uuid4()]):
                with pytest.raises(WorkspaceRepositoryError) as denied:
                    await photo_roster_service.source(connection, employee,
                        PhotoRosterExport(user_ids=identities), storage)
                assert denied.value.status_code == 403
            await dossier_service.save_access(connection, admin, DossierAccessWrite(
                revision=1, reason="QA revoke", rules=[],
            ))
            with pytest.raises(WorkspaceRepositoryError) as revoked:
                await photo_roster_service.finish(connection, employee, scoped, "pdf")
            assert revoked.value.status_code == 403
            count = (await connection.execute(select(audit_events.c.id).where(
                audit_events.c.action == "dossier.photo_roster_exported",
            ))).all()
            assert not count
    asyncio.run(run())


def test_failed_photo_and_pdf_do_not_generate_successful_export(monkeypatch):
    async def run():
        async with scenario() as (connection, actors):
            target, actor = actors["employee"], actors["deputy"]
            await connection.execute(update(users).where(users.c.id == target.id).values(
                avatar_storage_key="qa-missing-photo",
            ))
            storage = AsyncMock()
            storage.get.side_effect = ObjectStorageError("QA storage failure")
            app = FastAPI()
            app.include_router(routes.router, prefix="/api/v1")
            app.state.object_storage, app.state.settings = storage, object()
            app.dependency_overrides[require_user] = lambda: actor

            async def test_connection():
                yield connection

            app.dependency_overrides[get_connection] = test_connection
            payload = {"userIds": [str(target.id)]}
            async with AsyncClient(
                transport=ASGITransport(app=app), base_url="http://test",
            ) as client:
                route = "/api/v1/personnel-dossiers/photos/export/pdf"
                assert (await client.post(route, json=payload)).status_code == 503
                storage.get.side_effect = None
                storage.get.return_value = b"broken image"
                assert (await client.post(route, json=payload)).status_code == 422
                storage.get.return_value = portrait()
                monkeypatch.setattr(routes, "render_word_pdf", AsyncMock(
                    side_effect=HTTPException(503, "QA converter unavailable"),
                ))
                assert (await client.post(route, json=payload)).status_code == 503
            assert not (await connection.execute(select(audit_events.c.id).where(
                audit_events.c.action == "dossier.photo_roster_exported",
            ))).all()
    asyncio.run(run())
