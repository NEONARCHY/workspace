from datetime import date
from io import BytesIO

from openpyxl import Workbook

from yuksalish_api.hr_workbook import HrWorkbookError, read_tenure_workbook


def workbook_bytes() -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "сентябрь"
    worksheet.cell(10, 13).value = date(2026, 8, 31)
    worksheet.cell(13, 3).value = 1
    worksheet.cell(13, 4).value = "Саидова Саида"
    worksheet.cell(13, 8).value = "Главный специалист"
    worksheet.cell(13, 10).value = date(2022, 1, 15)
    worksheet.cell(13, 12).value = 4
    worksheet.cell(13, 13).value = 7
    worksheet.cell(13, 14).value = 16
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_read_tenure_workbook_returns_only_confirmed_rows() -> None:
    source_label, anchor, rows = read_tenure_workbook(workbook_bytes(), "2026.xlsx")

    assert source_label == "2026.xlsx:сентябрь:2026-08-31"
    assert anchor == date(2026, 8, 31)
    assert rows == [
        {
            "source_row": 13,
            "full_name": "Саидова Саида",
            "job_title": "Главный специалист",
            "employment_date": date(2022, 1, 15),
            "service_anchor_date": date(2026, 8, 31),
            "service_years": 4,
            "service_months": 7,
            "service_days": 16,
            "service_reason": "Импорт из 2026.xlsx, лист сентябрь, контрольная дата 2026-08-31",
        }
    ]


def test_read_tenure_workbook_rejects_wrong_extension() -> None:
    try:
        read_tenure_workbook(workbook_bytes(), "2026.xls")
    except HrWorkbookError as error:
        assert str(error) == "Поддерживается только файл Excel формата .xlsx"
    else:
        raise AssertionError("Expected HR workbook error")
