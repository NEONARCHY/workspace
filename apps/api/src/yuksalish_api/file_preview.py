"""Read-only, bounded Office content; no HTML, macros or external resources."""

from io import BytesIO
from itertools import islice
from pathlib import PurePath
from typing import Literal
from xml.etree.ElementTree import Element
from zipfile import ZipFile

import openpyxl  # type: ignore[import-untyped]
from defusedxml import ElementTree
from pydantic import Field

from .project_import_documents import MAX_TEXT_CHARS, document_type
from .workspace_schemas import ApiModel


class PreviewBlock(ApiModel):
    kind: Literal["paragraph", "table"]
    text: str = ""
    rows: list[list[str]] = Field(default_factory=list)


class OfficePreview(ApiModel):
    kind: Literal["word", "excel"]
    blocks: list[PreviewBlock] = Field(default_factory=list)
    sheets: list[str] = Field(default_factory=list)
    sheet_index: int = 0
    row_offset: int = 0
    total_rows: int = 0
    rows: list[list[str]] = Field(default_factory=list)
    truncated: bool = False


def preview_office(name: str, content: bytes, sheet_index: int, row_offset: int) -> OfficePreview:
    document_type(name, content)  # Signature, zip-bomb, macro and external-link checks.
    if PurePath(name).suffix.lower() == ".docx":
        ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        with ZipFile(BytesIO(content)) as archive:
            if archive.getinfo("word/document.xml").file_size > 4_000_000:
                raise ValueError("Word document exceeds preview limits")
            root = ElementTree.fromstring(archive.read("word/document.xml"))
        body = root.find(ns + "body")
        if body is None:
            raise ValueError("Документ не содержит текста")
        blocks: list[PreviewBlock] = []
        budget = MAX_TEXT_CHARS
        truncated = False

        def text(node: Element) -> str:
            return "".join(part.text or "" for part in node.iter(ns + "t"))

        for node in body:
            if len(blocks) >= 2000 or budget <= 0:
                truncated = True
                break
            if node.tag == ns + "p":
                value = text(node)
                blocks.append(PreviewBlock(kind="paragraph", text=value[:budget]))
                budget -= len(value)
            elif node.tag == ns + "tbl":
                rows: list[list[str]] = []
                for row in islice(node.findall(ns + "tr"), 200):
                    values: list[str] = []
                    cells = row.findall(ns + "tc")
                    truncated |= len(cells) > 80
                    for cell in islice(cells, 80):
                        value = text(cell)
                        shown = value[:min(2000, max(0, budget))]
                        truncated |= len(shown) < len(value)
                        values.append(shown)
                        budget -= len(shown)
                    rows.append(values)
                    if budget <= 0:
                        break
                truncated |= len(node.findall(ns + "tr")) > len(rows) or budget < 0
                blocks.append(PreviewBlock(kind="table", rows=rows))
        return OfficePreview(kind="word", blocks=blocks, truncated=truncated or budget < 0)
    if PurePath(name).suffix.lower() != ".xlsx":
        raise ValueError("Для просмотра сохраните документ в DOCX или XLSX")
    workbook = openpyxl.load_workbook(
        BytesIO(content), read_only=True, data_only=True, keep_links=False,
    )
    formulas = None
    try:
        if not 0 <= sheet_index < len(workbook.worksheets):
            raise ValueError("Лист не найден")
        sheet = workbook.worksheets[sheet_index]
        total_rows = sheet.max_row or 0
        total_columns = sheet.max_column or 0
        if total_rows * total_columns > 150_000:
            raise ValueError("Лист слишком большой для предпросмотра; скачайте оригинал")
        rows = [
            [str(value)[:2000] if value is not None else "" for value in row]
            for row in sheet.iter_rows(
                min_row=row_offset + 1, max_row=min(total_rows, row_offset + 200),
                max_col=min(total_columns, 80), values_only=True,
            )
        ] if total_rows and row_offset < total_rows else []
        # Cached results are displayed when present; uncached formulas remain
        # readable text and are never executed by the viewer.
        if rows:
            formulas = openpyxl.load_workbook(
                BytesIO(content), read_only=True, data_only=False, keep_links=False,
            )
            raw_rows = formulas.worksheets[sheet_index].iter_rows(
                min_row=row_offset + 1, max_row=min(total_rows, row_offset + 200),
                max_col=min(total_columns, 80), values_only=True,
            )
            for shown_row, raw in zip(rows, raw_rows, strict=True):
                for index, value in enumerate(raw):
                    if not shown_row[index] and isinstance(value, str) and value.startswith("="):
                        shown_row[index] = value[:2000]
        if sum(len(value) for row in rows for value in row) > MAX_TEXT_CHARS:
            raise ValueError("Слишком много текста на странице; скачайте оригинал")
        return OfficePreview(
            kind="excel", sheets=workbook.sheetnames, sheet_index=sheet_index,
            row_offset=row_offset, total_rows=total_rows, rows=rows,
            truncated=total_columns > 80,
        )
    finally:
        workbook.close()
        if formulas is not None:
            formulas.close()
