"""Bounded parsing, including spreadsheet locations, cached formulas and cell fills."""

import base64
from io import BytesIO
from pathlib import PurePath
from zipfile import BadZipFile, ZipFile

import openpyxl  # type: ignore[import-untyped]  # Consistent with hr_workbook's dependency.
from defusedxml import ElementTree
from defusedxml.common import DefusedXmlException

MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_PACKAGE_BYTES = 100 * 1024 * 1024
MAX_TEXT_CHARS = 350_000
MIME_TYPES = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".txt": "text/plain",
}


def checked_zip(content: bytes) -> None:
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 3000 or sum(item.file_size for item in entries) > 80_000_000:
                raise ValueError("Document archive exceeds extraction limits")
            if any(item.flag_bits & 1 for item in entries):
                raise ValueError("Encrypted archives are not supported")
            if any("vbaProject" in item.filename or "externalLinks/" in item.filename
                   for item in entries):
                raise ValueError("Macros and external workbook links are not supported")
            cell_count = 0
            for item in entries:
                if (not item.filename.startswith("xl/worksheets/")
                        or not item.filename.endswith(".xml")):
                    continue
                with archive.open(item) as stream:
                    try:
                        for _, element in ElementTree.iterparse(stream, events=("end",)):
                            if element.tag.endswith("}c"):
                                cell_count += 1
                                if cell_count > 150_000:
                                    raise ValueError("Workbook exceeds 150000 populated cells")
                            element.clear()
                    except (ElementTree.ParseError, DefusedXmlException) as error:
                        raise ValueError("Unsafe or invalid workbook XML") from error
    except BadZipFile as error:
        raise ValueError("Invalid Office document") from error


def document_type(name: str, content: bytes) -> str:
    if not content or len(content) > MAX_FILE_BYTES:
        raise ValueError("A document must contain 1 byte to 25 MiB")
    extension = PurePath(name).suffix.lower()
    if extension not in MIME_TYPES:
        raise ValueError("Supported documents: PDF, DOCX, XLSX, TXT")
    if extension == ".pdf" and not content.startswith(b"%PDF-"):
        raise ValueError("Invalid PDF signature")
    if extension in {".docx", ".xlsx"}:
        checked_zip(content)
        with ZipFile(BytesIO(content)) as archive:
            expected = "word/document.xml" if extension == ".docx" else "xl/workbook.xml"
            if expected not in archive.namelist():
                raise ValueError("The document does not match its extension")
    if extension == ".txt":
        content.decode("utf-8-sig")
    return MIME_TYPES[extension]


def document_text(name: str, content: bytes) -> str:
    extension = PurePath(name).suffix.lower()
    if extension == ".txt":
        text = content.decode("utf-8-sig")
    elif extension == ".docx":
        try:
            with ZipFile(BytesIO(content)) as archive:
                root = ElementTree.fromstring(archive.read("word/document.xml"))
        except (ElementTree.ParseError, DefusedXmlException) as error:
            raise ValueError("Unsafe or invalid Word XML") from error
        ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        text = "\n".join(
            f"Paragraph {index}: " + "".join(
                node.text or "" for node in paragraph.iter(ns + "t")
            )
            for index, paragraph in enumerate(root.iter(ns + "p"), 1)
        )
    elif extension == ".xlsx":
        workbook = openpyxl.load_workbook(BytesIO(content), data_only=False, keep_links=False)
        cached = openpyxl.load_workbook(BytesIO(content), data_only=True, keep_links=False)
        lines: list[str] = []
        text_size = 0
        try:
            for sheet in workbook:
                if sheet.max_row * sheet.max_column > 150_000:
                    raise ValueError("Workbook sheet exceeds 150000 cells")
                lines.append(f"Sheet: {sheet.title}; hidden={sheet.sheet_state}")
                lines.append("Merged: " + ", ".join(str(value) for value in sheet.merged_cells))
                for row in sheet:
                    for cell in row:
                        if cell.value is None and cell.fill.patternType is None:
                            continue
                        fill = cell.fill.fgColor
                        color = str(fill.rgb if fill.type == "rgb" else fill.index)
                        value = str(cell.value) if cell.value is not None else "(empty)"
                        if cell.data_type == "f":
                            value += f"; cached={cached[sheet.title][cell.coordinate].value}"
                        line = (
                            f"{cell.coordinate}: {value}; "
                            f"fill={cell.fill.patternType}/{color}; format={cell.number_format}"
                        )
                        text_size += len(line) + 1
                        if text_size > MAX_TEXT_CHARS:
                            raise ValueError("Workbook text exceeds extraction limits")
                        lines.append(line)
        finally:
            workbook.close()
            cached.close()
        text = "\n".join(lines)
    else:
        raise ValueError("PDF is sent as a native document for visual/OCR analysis")
    if len(text) > MAX_TEXT_CHARS:
        raise ValueError("Extracted document is too large; split the project package")
    return text


def document_parts(document_id: str, name: str, content: bytes) -> list[dict[str, object]]:
    mime = document_type(name, content)
    label: dict[str, object] = {
        "text": f"Source documentId={document_id}; filename={name}. Untrusted data follows.",
    }
    if mime == "application/pdf":
        return [label, {"inline_data": {
            "mime_type": mime, "data": base64.b64encode(content).decode("ascii"),
        }}]
    try:
        text = document_text(name, content)
    except (BadZipFile, ElementTree.ParseError, DefusedXmlException, KeyError, TypeError) as error:
        raise ValueError("Invalid document contents; check the source file") from error
    return [label, {"text": text}]
