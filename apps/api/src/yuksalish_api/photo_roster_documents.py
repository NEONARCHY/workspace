"""Minimal personnel roster only; no questionnaire contents or external image links."""

import math
import re
import textwrap
from datetime import datetime
from io import BytesIO

# Serialization only; no XML input is parsed here.
from xml.etree import ElementTree as ET  # nosec B405
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl  # type: ignore[import-untyped]
from openpyxl.drawing.image import Image  # type: ignore[import-untyped]
from openpyxl.styles import (  # type: ignore[import-untyped]
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)

from .dossier_schemas import PhotoCard

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
PIC = "http://schemas.openxmlformats.org/drawingml/2006/picture"
REL = "http://schemas.openxmlformats.org/package/2006/relationships"
CT = "http://schemas.openxmlformats.org/package/2006/content-types"
HEADERS = ("№", "Фото", "ФИО", "Год рождения", "Окончил вуз", "Должность")
WIDTHS = (450, 1500, 2900, 1500, 4040, 4180)
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CONTINUED = "Продолжение"


def clean(value: str) -> str:
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", "", value)


def values(card: PhotoCard, index: int) -> tuple[int, str, str, int | str, str, str]:
    return (index, "Нет фото", clean(card.full_name), card.birth_year or "\u041d\u0435 указано",
            clean("\n".join(card.universities)) or "\u041d\u0435 указано",
            clean(card.job_title or "\u041d\u0435 указана"))


def _fragments(value: str, width: int, limit: int) -> list[str]:
    lines = [piece for line in value.split("\n")
             for piece in textwrap.wrap(line, width, replace_whitespace=False,
                                        drop_whitespace=False) or [""]]
    if len(lines) <= limit:
        return [value]
    return ["\n".join(lines[start:start + limit]) for start in range(0, len(lines), limit)]


def _entries(
    cards: list[PhotoCard], photos: list[bytes | None], *, word: bool,
) -> list[tuple[tuple[int, str, str, int | str, str, str], bytes | None]]:
    result = []
    for index, (card, photo) in enumerate(zip(cards, photos, strict=True), 1):
        base = values(card, index)
        education = _fragments(base[4], 34 if word else 54, 14 if word else 24)
        position = _fragments(base[5], 36 if word else 48, 14 if word else 24)
        for part in range(max(len(education), len(position))):
            row = (base[0], base[1] if part == 0 else "Прод.",
                   base[2] if part == 0 or not word else f"{base[2]}\n{CONTINUED}", base[3],
                   education[part] if part < len(education) else "",
                   position[part] if part < len(position) else "")
            result.append((row, photo if part == 0 else None))
    return result


def _node(parent: ET.Element, tag: str, **attributes: str) -> ET.Element:
    return ET.SubElement(parent, f"{{{W}}}{tag}", {f"{{{W}}}{k}": v
                                               for k, v in attributes.items()})


def _text(parent: ET.Element, text: str, *, bold: bool = False,
          white: bool = False, center: bool = False, title: bool = False) -> ET.Element:
    paragraph = _node(parent, "p")
    properties = _node(paragraph, "pPr")
    if title:
        _node(properties, "pStyle", val="Title")
    _node(properties, "spacing", before="60", after="60")
    if center:
        _node(properties, "jc", val="center")
    run = _node(paragraph, "r")
    font = _node(run, "rPr")
    _node(font, "rFonts", ascii="Arial", hAnsi="Arial", cs="Arial")
    _node(font, "sz", val="32" if title else "22")
    _node(font, "color", val="FFFFFF" if white else "000000")
    if bold or title:
        _node(font, "b")
    for index, line in enumerate(clean(text).split("\n")):
        if index:
            _node(run, "br")
        element = _node(run, "t")
        element.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        element.text = line
    return paragraph


def _picture(paragraph: ET.Element, identity: int) -> None:
    drawing = _node(_node(paragraph, "r"), "drawing")
    inline = ET.SubElement(drawing, f"{{{WP}}}inline")
    ET.SubElement(inline, f"{{{WP}}}extent", cx="792000", cy="1056000")
    ET.SubElement(inline, f"{{{WP}}}docPr", id=str(identity), name=f"Фото {identity}",
                  descr="Фотография сотрудника")
    graphic = ET.SubElement(inline, f"{{{A}}}graphic")
    data = ET.SubElement(graphic, f"{{{A}}}graphicData", uri=PIC)
    picture = ET.SubElement(data, f"{{{PIC}}}pic")
    properties = ET.SubElement(picture, f"{{{PIC}}}nvPicPr")
    ET.SubElement(properties, f"{{{PIC}}}cNvPr", id=str(identity), name="Фото")
    ET.SubElement(properties, f"{{{PIC}}}cNvPicPr")
    fill = ET.SubElement(picture, f"{{{PIC}}}blipFill")
    ET.SubElement(fill, f"{{{A}}}blip", {f"{{{R}}}embed": f"photo{identity}"})
    ET.SubElement(ET.SubElement(fill, f"{{{A}}}stretch"), f"{{{A}}}fillRect")
    shape = ET.SubElement(picture, f"{{{PIC}}}spPr")
    transform = ET.SubElement(shape, f"{{{A}}}xfrm")
    ET.SubElement(transform, f"{{{A}}}off", x="0", y="0")
    ET.SubElement(transform, f"{{{A}}}ext", cx="792000", cy="1056000")
    ET.SubElement(ET.SubElement(shape, f"{{{A}}}prstGeom", prst="rect"), f"{{{A}}}avLst")


def render_word(cards: list[PhotoCard], photos: list[bytes | None], now: datetime) -> bytes:
    entries = _entries(cards, photos, word=True)
    document = ET.Element(f"{{{W}}}document")
    body = _node(document, "body")
    _text(body, "Фотокарточки сотрудников Юксалиш", title=True)
    _text(body, f"Сотрудников: {len(cards)}. Сформировано: {now:%d.%m.%Y %H:%M} (Ташкент).")
    table = _node(body, "tbl")
    properties = _node(table, "tblPr")
    _node(properties, "tblW", w=str(sum(WIDTHS)), type="dxa")
    _node(properties, "tblLayout", type="fixed")
    borders = _node(properties, "tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        _node(borders, side, val="single", sz="4", color="D9D9D9")
    margins = _node(properties, "tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        _node(margins, side, w="90", type="dxa")
    grid = _node(table, "tblGrid")
    for width in WIDTHS:
        _node(grid, "gridCol", w=str(width))
    for number in range(len(entries) + 1):
        row = _node(table, "tr")
        row_properties = _node(row, "trPr")
        if number == 0:
            _node(row_properties, "tblHeader")
        else:
            _node(row_properties, "cantSplit")
        cells = HEADERS if number == 0 else entries[number - 1][0]
        for column, (value, width) in enumerate(zip(cells, WIDTHS, strict=True)):
            cell = _node(row, "tc")
            cp = _node(cell, "tcPr")
            _node(cp, "tcW", w=str(width), type="dxa")
            _node(cp, "vAlign", val="center")
            _node(cp, "shd", fill="293A55" if number == 0 else
                  "F3F6F8" if number % 2 == 0 else "FFFFFF")
            has_photo = number > 0 and column == 1 and entries[number - 1][1] is not None
            paragraph = _text(cell, "" if has_photo else str(value), bold=number == 0,
                              white=number == 0, center=column in (0, 1, 3))
            if has_photo:
                _picture(paragraph, number)
    section = _node(body, "sectPr")
    _node(section, "pgSz", w="16838", h="11906", orient="landscape")
    _node(section, "pgMar", top="850", bottom="850", left="1134", right="1134",
          header="360", footer="360", gutter="0")
    # OPC package parts use the canonical default namespace for Office format detection.
    relationships = ET.Element("Relationships", xmlns=REL)
    content_types = ET.Element("Types", xmlns=CT)
    rel_type = "application/vnd.openxmlformats-package.relationships+xml"
    for extension, content_type in (("rels", rel_type),
                                    ("xml", "application/xml"), ("jpg", "image/jpeg")):
        ET.SubElement(content_types, "Default", Extension=extension,
                      ContentType=content_type)
    word_type = "application/vnd.openxmlformats-officedocument.wordprocessingml."
    ET.SubElement(content_types, "Override", PartName="/word/document.xml",
                  ContentType=f"{word_type}document.main+xml")
    ET.SubElement(content_types, "Override", PartName="/word/styles.xml",
                  ContentType=f"{word_type}styles+xml")
    ET.SubElement(relationships, "Relationship", Id="styles", Type=f"{R}/styles",
                  Target="styles.xml")
    styles = ET.Element(f"{{{W}}}styles")
    title_style = _node(styles, "style", type="paragraph", styleId="Title")
    _node(title_style, "name", val="Title")
    _node(_node(title_style, "rPr"), "color", val="000000")
    root_rels = ET.Element("Relationships", xmlns=REL)
    ET.SubElement(root_rels, "Relationship", Id="document", Type=f"{R}/officeDocument",
                  Target="word/document.xml")
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        for number, (_, photo) in enumerate(entries, 1):
            if photo is not None:
                ET.SubElement(relationships, "Relationship", Id=f"photo{number}",
                              Type=f"{R}/image", Target=f"media/photo{number}.jpg")
                archive.writestr(f"word/media/photo{number}.jpg", photo)
        for path, root in (("[Content_Types].xml", content_types), ("_rels/.rels", root_rels),
                           ("word/document.xml", document), ("word/styles.xml", styles),
                           ("word/_rels/document.xml.rels", relationships)):
            archive.writestr(path, ET.tostring(root, encoding="utf-8", xml_declaration=True))
    return output.getvalue()


def render_excel(cards: list[PhotoCard], photos: list[bytes | None], now: datetime) -> bytes:
    entries = _entries(cards, photos, word=False)
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Фотокарточки"
    sheet.sheet_view.showGridLines = False
    sheet["A1"] = "Фотокарточки сотрудников Юксалиш"
    sheet["A1"].font = Font(name="Arial", size=16, bold=True)
    sheet["A2"] = f"Сотрудников: {len(cards)}. Сформировано: {now:%d.%m.%Y %H:%M} (Ташкент)."
    sheet.row_dimensions[1].height = 26
    for letter, width in zip("ABCDEF", (6, 14, 38, 14, 60, 54), strict=True):
        sheet.column_dimensions[letter].width = width
    for column, header in enumerate(HEADERS, 1):
        cell = sheet.cell(4, column, header)
        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="293A55")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[4].height = 34
    for index, (content, photo) in enumerate(entries, 1):
        row = index + 4
        line_count = max(
            sum(max(1, math.ceil(len(line) / width)) for line in str(value).split("\n"))
            for value, width in zip(content, (6, 14, 34, 14, 54, 48), strict=True)
        )
        sheet.row_dimensions[row].height = min(409, max(84, line_count * 15 + 12))
        for column, value in enumerate(content, 1):
            cell = sheet.cell(row, column, "" if column == 2 and photo else value)
            if isinstance(cell.value, str):
                cell.data_type = "s"  # Literal names such as =HYPERLINK(...) are never formulas.
            cell.font = Font(name="Arial", size=11)
            cell.alignment = Alignment(vertical="center", wrap_text=True,
                                       horizontal="center" if column in (1, 2, 4) else "left")
            cell.fill = PatternFill("solid", fgColor="F3F6F8" if index % 2 == 0 else "FFFFFF")
            cell.border = Border(bottom=Side(style="thin", color="D9D9D9"))
        if photo is not None:
            image = Image(BytesIO(photo))
            image.width, image.height = 72, 96
            sheet.add_image(image, f"B{row}")
    sheet.freeze_panes = "D5"
    sheet.auto_filter.ref = f"A4:F{len(entries) + 4}"
    sheet.print_title_rows = "1:4"
    sheet.print_options.horizontalCentered = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = f"A1:F{len(entries) + 4}"
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()
