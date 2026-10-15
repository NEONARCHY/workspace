"""Read-only extraction of plan candidates and reviewable table-cell hints."""

from __future__ import annotations

import re
from calendar import monthrange
from collections.abc import Iterator
from dataclasses import dataclass
from datetime import date
from io import BytesIO
from pathlib import PurePath
from typing import TYPE_CHECKING, Protocol
from zipfile import ZipFile

import openpyxl  # type: ignore[import-untyped]
from defusedxml import ElementTree
from pypdf import PdfReader

from .project_import_documents import checked_zip, document_type
from .task_plan_schemas import MAX_PLAN_IMPORT_ITEMS, TaskPlanImportPreview, TaskPlanImportRow

if TYPE_CHECKING:
    # Type only: document bytes are parsed exclusively by defusedxml below.
    from xml.etree.ElementTree import Element  # nosec B405

MAX_PLAN_BYTES = 10 * 1024 * 1024
MAX_ROWS = MAX_PLAN_IMPORT_ITEMS
MAX_TEXT_CHARS = 50_000
_NUMBER = re.compile(r"^\s*(?:\d+[.)]|\d+(?:\.\d+)+[.)]?|[•●▪◦‣-])\s*")
_TABLE_BULLET = re.compile(r"^\s*(?:[•●▪◦‣]\s*|-\s+)")
_PERSON_HEADING = re.compile(r"\b[^\W\d_]\s*\.\s*[^\W\d_][^\W\d_'-]+", re.UNICODE)
_DATE_CELL = re.compile(r"^\s*(?:до\s+)?(\d{1,2})[./](\d{1,2})[./](\d{4})\s*$")
_WEEK_END = re.compile(r"^\s*(\d{1,2})(?:\s+([\u0430-\u044f\u0451]+))?\s*"
                       r"[\u2013\u2014-]\s*(\d{1,2})\s+([\u0430-\u044f\u0451]+)\s*$",
                       re.IGNORECASE)
_PLAN_YEAR = re.compile(r"\b20\d{2}\b")
_MONTHS = {
    "январь": 1, "января": 1, "февраль": 2, "февраля": 2,
    "март": 3, "марта": 3, "апрель": 4, "апреля": 4,
    "май": 5, "мая": 5, "июнь": 6, "июня": 6,
    "июль": 7, "июля": 7, "август": 8, "августа": 8,
    "сентябрь": 9, "сентября": 9, "октябрь": 10, "октября": 10,
    "ноябрь": 11, "ноября": 11, "декабрь": 12, "декабря": 12,
}
_ANNUAL_MONTH = re.compile(
    r"январ[ья]?|феврал[ья]?|март[а]?|апрел[ья]?|ма[йя]|июн[ья]?|"  # noqa: RUF001
    r"июл[ья]?|август[а]?|сентябр[ья]?|октябр[ья]?|ноябр[ья]?|декабр[ья]?",  # noqa: RUF001
    re.IGNORECASE,
)
_ANNUAL_MONTH_NUMBER = {
    "янв": 1, "фев": 2, "мар": 3, "апр": 4, "май": 5, "мая": 5,
    "июн": 6, "июл": 7, "авг": 8, "сен": 9, "окт": 10, "ноя": 11, "дек": 12,
}
_ANNUAL_RECURRING = re.compile(
    r"ҳар\s+(?:ой|чорак|хафта)|йил\s+давомида|йилда\s+\d|режа-график",
    re.IGNORECASE,
)
_ANNUAL_MOVED = re.compile(r"кўчирил|ўзгартирил|ўзгарди", re.IGNORECASE)


@dataclass(frozen=True)
class _Candidate:
    title: str
    source: str
    assignee_hint: str | None = None
    coexecutors_hint: str | None = None
    group_key: str | None = None
    due_date_hint: date | None = None
    deadline_text: str | None = None
    mechanism_items: tuple[str, ...] = ()
    completed_work_text: str | None = None
    original_title: str | None = None
    responsible_text: str | None = None
    executors_text: str | None = None
    section: str | None = None
    green_completion: bool = False
    deadline_needs_review: bool = False
    warning_hint: str | None = None
    annual_plan: bool = False
    plan_year: int | None = None


class _XmlNode(Protocol):
    tag: str
    text: str | None

    def iter(self, tag: str | None = None) -> Iterator[_XmlNode]: ...


def _docx_paragraph_text(paragraph: _XmlNode, ns: str) -> str:
    separators = {ns + "br", ns + "cr", ns + "tab"}
    return "".join(
        " " if node.tag in separators else node.text or ""
        for node in paragraph.iter() if node.tag == ns + "t" or node.tag in separators
    )


def _cell_text(cell: _XmlNode, ns: str) -> str:
    return " ".join(" ".join(_docx_paragraph_text(paragraph, ns).split())
                    for paragraph in cell.iter(ns + "p")).strip()


def _cell_lines(cell: _XmlNode, ns: str) -> list[str]:
    return [line for paragraph in cell.iter(ns + "p")
            if (line := " ".join(_docx_paragraph_text(paragraph, ns).split()))]


def _annual_due_hint(value: str, year: int | None) -> tuple[date | None, bool]:
    """Suggest the last stated deadline, never pretending a recurring rule is a date."""
    exact = _date_cell_hint(value)
    if exact is not None:
        return exact, False
    if year is None or _ANNUAL_RECURRING.search(value):
        return None, True
    lowered = value.casefold()
    if "биринчи ярми" in lowered:
        return date(year, 6, 30), True
    months = list(_ANNUAL_MONTH.finditer(lowered))
    if not months:
        if re.search(r"\biv\s+чорак", lowered):
            return date(year, 12, 31), True
        return None, True
    last = months[-1]
    month = _ANNUAL_MONTH_NUMBER[last.group()[:3]]
    before = lowered[:last.start()]
    days = re.search(r"(?<!\d)(\d{1,2})(?:\s*[-–—]\s*(\d{1,2}))?\s*[- ]*$", before)  # noqa: RUF001
    day = int(days.group(2) or days.group(1)) if days else monthrange(year, month)[1]
    try:
        due = date(year, month, day)
    except ValueError:
        return None, True
    return due, bool(_ANNUAL_MOVED.search(value) or len(months) > 1 or "," in value)


def _green_completion(cell: Element, ns: str) -> bool:
    colors: dict[str, int] = {}
    total = 0
    for run in cell.iter(ns + "r"):
        text = "".join(node.text or "" for node in run.iter(ns + "t")).strip()
        if not text:
            continue
        total += len(text)
        color = run.find(f"{ns}rPr/{ns}color")
        if color is not None:
            value = (color.get(ns + "val") or "").upper()
            colors[value] = colors.get(value, 0) + len(text)
    body = _cell_text(cell, ns).casefold()
    return bool(total and colors.get("00B050", 0) >= total * 0.8
                and not any(colors.get(value, 0) for value in
                            ("C00000", "FF9900", "E69138", "FFFF00"))
                and not any(word in body for word in
                            ("бажарилмади", "жараёнда", "муддати бор")))  # noqa: RUF001


def _annual_docx_rows(body: Element, ns: str, plan_year: int | None) -> list[_Candidate] | None:
    """Read the seven semantic columns in the 2026 work-plan table, including merged headers."""
    result: list[_Candidate] = []
    found = False
    for table_number, table in enumerate(body.findall(ns + "tbl"), 1):
        table_rows = table.findall(ns + "tr")
        if not table_rows:
            continue
        header = [_cell_text(cell, ns).casefold() for cell in table_rows[0].findall(ns + "tc")]
        if len(header) != 7 or not all(any(token in cell for cell in header) for token in (
            "тадбирнинг номи", "муддати", "амалга ошириш механизми",
            "бажарувчи", "масъул", "бажарилган ишлар",
        )):
            continue
        found = True
        section: str | None = None
        for row_number, row in enumerate(table_rows[1:], 2):
            cells = row.findall(ns + "tc")
            if len(cells) == 1:
                section = _cell_text(cells[0], ns)[:500] or section
                continue
            if len(cells) != 7:
                continue
            title = _cell_text(cells[1], ns)
            if len(title) < 5:
                continue
            deadline = _cell_text(cells[2], ns)
            due_date, needs_review = _annual_due_hint(deadline, plan_year)
            mechanism = tuple(_NUMBER.sub("", line).strip()
                              for line in _cell_lines(cells[3], ns))
            mechanism = tuple(line for line in mechanism if line)
            executors = _cell_text(cells[4], ns)
            responsible = _cell_text(cells[5], ns)
            first_executor = _PERSON_HEADING.search(executors)
            executor_hint = (first_executor.group() if first_executor else
                             next((part.strip() for part in re.split(r"[;,\n]", executors)
                                   if part.strip()), None))
            completed_work = "\n\n".join(_cell_lines(cells[6], ns))
            warnings = []
            if len(title) > 240:
                warnings.append(
                    "Название длиннее поля задачи; полный оригинал сохранится в описании"
                )
            if not executors:
                warnings.append("В колонке «Бажарувчи» исполнители не указаны")  # noqa: RUF001
            result.append(_Candidate(
                title=title[:239] + "…" if len(title) > 240 else title,
                source=f"DOCX, таблица {table_number}, строка {row_number}",
                assignee_hint=executor_hint,
                coexecutors_hint=executors or None,
                due_date_hint=due_date,
                deadline_text=deadline or None,
                mechanism_items=mechanism,
                completed_work_text=completed_work or None,
                original_title=title if len(title) > 240 else None,
                responsible_text=responsible or None,
                executors_text=executors or None,
                section=section,
                green_completion=_green_completion(cells[6], ns),
                deadline_needs_review=needs_review,
                warning_hint="; ".join(warnings) or None,
                annual_plan=True,
                plan_year=plan_year,
            ))
            if len(result) > MAX_ROWS:
                return result
    return result if found else None


def _date_cell_hint(value: str, plan_year: int | None = None) -> date | None:
    match = _DATE_CELL.fullmatch(value)
    if match is not None:
        try:
            return date(int(match[3]), int(match[2]), int(match[1]))
        except ValueError:
            return None
    week = _WEEK_END.fullmatch(value)
    if week is None or plan_year is None:
        return None
    month = _MONTHS.get(week[4].casefold())
    if month is None:
        return None
    start_month = _MONTHS.get(week[2].casefold()) if week[2] else month
    if start_month is None or start_month > month:
        return None
    try:
        start = date(plan_year, start_month, int(week[1]))
        end = date(plan_year, month, int(week[3]))
    except ValueError:
        return None
    return end if start <= end else None


def _text_rows(text: str, location: str) -> list[_Candidate]:
    if len(text) > MAX_TEXT_CHARS:
        raise ValueError("Document text exceeds extraction limits")
    rows: list[_Candidate] = []
    for number, line in enumerate(text.splitlines(), 1):
        title = _NUMBER.sub("", " ".join(line.split())).strip()
        if len(title) < 5 or title.lower() in {"план работ", "иш режа", "мероприятия"}:
            continue
        rows.append(_Candidate(title[:240], f"{location}, строка {number}"))
        if len(rows) > MAX_ROWS:
            break
    return rows


def _docx_rows(content: bytes) -> list[_Candidate]:
    with ZipFile(BytesIO(content)) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    body = root.find(ns + "body")
    if body is not None:
        heading_text = " ".join(_docx_paragraph_text(paragraph, ns)
                                for paragraph in body.findall(ns + "p"))
        years = set(_PLAN_YEAR.findall(heading_text))
        plan_year = int(next(iter(years))) if len(years) == 1 else None
        annual_rows = _annual_docx_rows(body, ns, plan_year)
        if annual_rows is not None:
            return annual_rows
    paragraphs = [_docx_paragraph_text(paragraph, ns) for paragraph in root.iter(ns + "p")]
    if sum(len(paragraph) for paragraph in paragraphs) > MAX_TEXT_CHARS:
        raise ValueError("Document text exceeds extraction limits")

    # In weekly-plan tables, each cell can contain several real tasks. Do not turn
    # the week, names, co-executors and KPI cells into candidate tasks.
    table_rows: list[_Candidate] = []
    if body is not None:
        for table_number, table in enumerate(body.findall(ns + "tbl"), 1):
            header = table.find(ns + "tr")
            header_cells = ([_cell_text(cell, ns)
                             for cell in header.findall(ns + "tc")] if header is not None else [])
            coexecutors_column = next((index for index, value in enumerate(header_cells)
                                       if any(label in value.casefold() for label in
                                              ("ҳамижрочи", "hamijrochi", "соисполнител"))), None)
            responsible_column = next((index for index, value in enumerate(header_cells)
                                       if any(label in value.casefold() for label in
                                              ("ответствен", "масъул", "mas'ul"))), None)
            deadline_column = next((index for index, value in enumerate(header_cells)
                                    if any(label in value.casefold() for label in
                                           ("срок", "муддат", "deadline"))), None)
            for row_number, row in enumerate(table.findall(ns + "tr"), 1):
                if row_number == 1:
                    continue
                cells = row.findall(ns + "tc")
                coexecutors_hint = (_cell_text(cells[coexecutors_column], ns)[:1000]
                                    if coexecutors_column is not None
                                    and coexecutors_column < len(cells) else "")
                row_assignee = (_cell_text(cells[responsible_column], ns)[:240]
                                if responsible_column is not None
                                and responsible_column < len(cells) else "")
                due_date_hint = (_date_cell_hint(_cell_text(cells[deadline_column], ns), plan_year)
                                 if deadline_column is not None
                                 and deadline_column < len(cells) else None)
                for column_number, cell in enumerate(cells, 1):
                    index = column_number - 1
                    if index in {coexecutors_column, responsible_column, deadline_column}:
                        continue
                    if coexecutors_column is not None and index >= coexecutors_column:
                        continue
                    heading_matches = (_PERSON_HEADING.findall(header_cells[index])
                                       if index < len(header_cells) else [])
                    assignee_hint = row_assignee or (
                        heading_matches[0] if len(heading_matches) == 1 else None
                    )
                    structured = (coexecutors_column is not None or responsible_column is not None
                                  or deadline_column is not None or bool(heading_matches))
                    group_key = (f"DOCX, таблица {table_number}, строка {row_number}"
                                 if structured else None)
                    for paragraph in cell.iter(ns + "p"):
                        line = _docx_paragraph_text(paragraph, ns)
                        if not _TABLE_BULLET.match(line):
                            continue
                        title = _NUMBER.sub("", " ".join(line.split())).strip()
                        if len(title) < 5:
                            continue
                        source = (f"DOCX, таблица {table_number}, строка {row_number}, "
                                  f"столбец {column_number}")
                        table_rows.append(_Candidate(
                            title[:240], source, assignee_hint,
                            coexecutors_hint or None, group_key, due_date_hint,
                        ))
                        if len(table_rows) > MAX_ROWS:
                            return table_rows
    if table_rows:
        return table_rows
    return _text_rows("\n".join(paragraphs), "DOCX")


def _xlsx_rows(content: bytes) -> list[_Candidate]:
    workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True,
                                      keep_links=False)
    rows: list[_Candidate] = []
    try:
        for sheet in workbook:
            if sheet.sheet_state != "visible":
                continue
            if sheet.max_row > 500 or sheet.max_column > 30:
                raise ValueError("Workbook exceeds 500 rows or 30 columns per sheet")
            for row_number, cells in enumerate(sheet.iter_rows(max_row=500, max_col=30), 1):
                texts = [" ".join(str(cell.value).split()) for cell in cells
                         if isinstance(cell.value, str) and len(cell.value.strip()) >= 5]
                if not texts:
                    continue
                # The longest text cell is generally the activity, not its number/budget.
                title = _NUMBER.sub("", max(texts, key=len)).strip()
                if len(title) < 5 or title.casefold() in {"activity", "мероприятие", "задача"}:
                    continue
                rows.append(_Candidate(title[:240], f"{sheet.title}!{row_number}"))
                if len(rows) > MAX_ROWS:
                    return rows
    finally:
        workbook.close()
    return rows


def preview_plan_document(filename: str, content: bytes) -> TaskPlanImportPreview:
    if len(content) > MAX_PLAN_BYTES:
        raise ValueError("План не должен превышать 10 МБ")
    document_type(filename, content)
    suffix = PurePath(filename).suffix.lower()
    if suffix in {".docx", ".xlsx"}:
        checked_zip(content)
    if suffix == ".xlsx":
        candidates = _xlsx_rows(content)
    elif suffix == ".docx":
        candidates = _docx_rows(content)
    elif suffix == ".txt":
        candidates = _text_rows(content.decode("utf-8-sig"), "TXT")
    else:
        reader = PdfReader(BytesIO(content), strict=False)
        if reader.is_encrypted or len(reader.pages) > 20:
            raise ValueError("PDF защищён или содержит более 20 страниц")
        candidates = []
        total_chars = 0
        for index, page in enumerate(reader.pages, 1):
            extracted = page.extract_text() or ""
            total_chars += len(extracted)
            if total_chars > MAX_TEXT_CHARS:
                raise ValueError("PDF text exceeds extraction limits")
            candidates.extend(_text_rows(extracted, f"PDF, стр. {index}"))
            if len(candidates) > MAX_ROWS:
                break
    if len(candidates) > MAX_ROWS:
        raise ValueError(f"План содержит более {MAX_ROWS} пунктов. Разделите файл на части.")
    seen: set[str] = set()
    rows: list[TaskPlanImportRow] = []
    for candidate in candidates:
        key = candidate.title.casefold()
        warning = "Похожий пункт уже есть в файле — проверьте дубль" if key in seen else None
        if candidate.warning_hint:
            warning = "; ".join(filter(None, (warning, candidate.warning_hint)))
        seen.add(key)
        rows.append(TaskPlanImportRow(
            title=candidate.title, source=candidate.source, warning=warning,
            assignee_hint=candidate.assignee_hint,
            coexecutors_hint=candidate.coexecutors_hint,
            group_key=candidate.group_key,
            due_date_hint=candidate.due_date_hint,
            deadline_text=candidate.deadline_text,
            mechanism_items=list(candidate.mechanism_items),
            completed_work_text=candidate.completed_work_text,
            original_title=candidate.original_title,
            responsible_text=candidate.responsible_text,
            executors_text=candidate.executors_text,
            section=candidate.section,
            green_completion=candidate.green_completion,
            deadline_needs_review=candidate.deadline_needs_review,
        ))
    if not rows:
        raise ValueError("No plan text found. Scanned PDFs require OCR; add a plan manually")
    annual = bool(candidates[0].annual_plan)
    suggested_title = (f"Иш режа {candidates[0].plan_year}" if annual and
                       candidates[0].plan_year else PurePath(filename).stem[:240])
    return TaskPlanImportPreview(suggested_title=suggested_title, rows=rows,
                                 annual_work_plan=annual)
