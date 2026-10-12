"""Spreadsheet export as safe literal strings, including headers and metadata."""
# Russian user-facing messages intentionally contain Cyrillic.
# ruff: noqa: RUF001

import io
import json

import openpyxl  # type: ignore[import-untyped]
from openpyxl.utils.exceptions import IllegalCharacterError  # type: ignore[import-untyped]

from .errors import WorkspaceRepositoryError
from .survey_schemas import SurveyExport, SurveyWorkbook
from .survey_service import validate_configuration


def export_xlsx(workbook: SurveyWorkbook, payload: SurveyExport) -> bytes:
    validate_configuration(payload.configuration, workbook)
    config = payload.configuration
    sheet = workbook.sheets[config.sheet]
    if (len(set(payload.row_numbers)) != len(payload.row_numbers)
            or any(number <= config.header_row + 1 or number > len(sheet.rows)
                   for number in payload.row_numbers)):
        raise WorkspaceRepositoryError(422, "Некорректные номера исходных строк.")
    book = openpyxl.Workbook(write_only=True)
    answers = book.create_sheet("Answers")
    metadata = book.create_sheet("Methodology")
    width = max((len(row) for row in sheet.rows[config.header_row:]), default=0)

    def literal_row(values: list[str]) -> list[object]:
        result: list[object] = []
        for value in values:
            cell = openpyxl.cell.WriteOnlyCell(answers, value=value)
            cell.data_type = "s"  # Never infer formulas, including header cells.
            result.append(cell)
        return result

    try:
        headers = sheet.rows[config.header_row]
        answers.append(literal_row(["Source row"] + [
            headers[index] if index < len(headers) else f"Column {index + 1}"
            for index in range(width)
        ]))
        for number in payload.row_numbers:
            row = sheet.rows[number - 1]
            answers.append(literal_row([str(number)] + [
                row[index] if index < len(row) else "" for index in range(width)
            ]))
        for key, value in [
            ("Sheet", sheet.name), ("Header source row", str(config.header_row + 1)),
            ("Selected source rows", str(len(payload.row_numbers))),
            ("Configuration", json.dumps(config.model_dump(mode="json"), ensure_ascii=False)),
            ("Method", "Unweighted descriptive analysis; one source row is one assumed respondent. "
             "Question percentages exclude missing answers. Multiple choice is deduplicated per "
             "row and may total above 100%. No representativeness or significance claim."),
            ("Privacy", "Export includes identifiers if present. Review before sharing."),
            ("Safety", "All cells are text; formulas are not executed. Original is unchanged."),
        ]:
            metadata.append(literal_row([key, value]))
        output = io.BytesIO()
        book.save(output)
        return output.getvalue()
    except IllegalCharacterError as error:
        raise WorkspaceRepositoryError(422, "В ответах есть символы, недопустимые в XLSX. "
                                       "Используйте CSV или исправьте исходный файл.") from error
    finally:
        book.close()
