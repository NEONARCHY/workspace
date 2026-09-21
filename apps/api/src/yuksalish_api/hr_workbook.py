# ruff: noqa: RUF001
from __future__ import annotations

import re
from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook  # type: ignore[import-untyped]


class HrWorkbookError(ValueError):
    """The workbook cannot be used as a tenure import source."""


def _date(value: object, field: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        match = re.search(r"(\d{2})\.(\d{2})\.(\d{4})", value)
        if match:
            day, month, year = map(int, match.groups())
            return date(year, month, day)
    raise HrWorkbookError(f"{field}: ожидалась дата")


def read_tenure_workbook(
    content: bytes, filename: str
) -> tuple[str, date, list[dict[str, object]]]:
    """Read the approved September-layout workbook without persisting its contents."""
    if not filename.lower().endswith(".xlsx"):
        raise HrWorkbookError("Поддерживается только файл Excel формата .xlsx")
    if len(f"{filename}:сентябрь:2026-08-31") > 120:
        raise HrWorkbookError("Имя файла слишком длинное для импорта")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (OSError, ValueError, KeyError) as error:
        raise HrWorkbookError("Не удалось открыть файл Excel") from error
    sheet_name = "сентябрь"
    if sheet_name not in workbook.sheetnames:
        raise HrWorkbookError("В файле не найден лист «сентябрь»")
    sheet = workbook[sheet_name]
    anchor = _date(sheet.cell(10, 13).value, "Контрольная дата")
    rows: list[dict[str, object]] = []
    for source_row in range(13, sheet.max_row + 1):
        number = sheet.cell(source_row, 3).value
        full_name = sheet.cell(source_row, 4).value
        if not isinstance(number, int) or not isinstance(full_name, str) or not full_name.strip():
            continue
        employment_date = _date(sheet.cell(source_row, 10).value, f"Строка {source_row}")
        service = tuple(sheet.cell(source_row, column).value for column in (12, 13, 14))
        if not all(isinstance(value, int) for value in service):
            raise HrWorkbookError(
                f"Строка {source_row}: стаж должен быть указан годами, месяцами и днями"
            )
        years, months, days = service
        if years < 0 or not 0 <= months <= 11 or not 0 <= days <= 30:
            raise HrWorkbookError(f"Строка {source_row}: указано некорректное значение стажа")
        title = sheet.cell(source_row, 8).value
        rows.append(
            {
                "source_row": source_row,
                "full_name": full_name.strip(),
                "job_title": title.strip() if isinstance(title, str) and title.strip() else None,
                "employment_date": employment_date,
                "service_anchor_date": anchor,
                "service_years": years,
                "service_months": months,
                "service_days": days,
                "service_reason": (
                    f"Импорт из {filename}, лист {sheet_name}, "
                    f"контрольная дата {anchor.isoformat()}"
                ),
            }
        )
    if not rows:
        raise HrWorkbookError("В листе не найдены строки сотрудников")
    source_label = f"{filename}:{sheet_name}:{anchor.isoformat()}"
    return source_label, anchor, rows
