"""Bounded workbook ingestion. Never executes formulas or silently truncates data."""
# Russian user-facing messages intentionally contain Cyrillic.
# ruff: noqa: RUF001

import csv
import io
import math
import re
import struct
from datetime import date, datetime
from pathlib import PurePath
from typing import Literal
from zipfile import ZipFile

import openpyxl  # type: ignore[import-untyped]
import xlrd  # type: ignore[import-untyped]
from defusedxml import ElementTree
from defusedxml.common import DefusedXmlException
from openpyxl.utils.exceptions import InvalidFileException  # type: ignore[import-untyped]

from .project_import_documents import checked_zip
from .survey_schemas import SurveySheet, SurveyWorkbook

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 25_000
MAX_COLUMNS = 512
MAX_CELLS = 150_000
MAX_TEXT = 2_000_000


class SurveyDocumentError(ValueError):
    """Safe user-facing validation message, never a raw parser exception."""


def check_coordinates(content: bytes) -> None:
    with ZipFile(io.BytesIO(content)) as archive:
        for item in archive.infolist():
            if not item.filename.endswith(".xml"):
                continue
            with archive.open(item) as stream:
                for _, element in ElementTree.iterparse(stream, events=("end",)):
                    if item.filename.startswith("xl/worksheets/") and element.tag.endswith("}c"):
                        coordinate = element.attrib.get("r", "")
                        match = re.fullmatch(r"([A-Z]{1,3})([1-9][0-9]{0,5})", coordinate)
                        if not match:
                            raise SurveyDocumentError("Некорректная координата ячейки XLSX.")
                        column = 0
                        for letter in match[1]:
                            column = column * 26 + ord(letter) - ord("A") + 1
                        if column > MAX_COLUMNS or int(match[2]) > MAX_ROWS:
                            raise SurveyDocumentError("Ячейка XLSX вне пределов 25 000 × 512.")
                    element.clear()


def parse_workbook(
    filename: str, content: bytes, *, encoding: Literal["utf-8-sig", "cp1251"] = "utf-8-sig",
    delimiter: Literal[",", ";", "\t"] = ",",
) -> SurveyWorkbook:
    try:
        return _parse_workbook(filename, content, encoding=encoding, delimiter=delimiter)
    except SurveyDocumentError:
        raise
    except (ValueError, KeyError, OSError, EOFError, struct.error, xlrd.XLRDError,
            ElementTree.ParseError, DefusedXmlException, InvalidFileException) as error:
        raise SurveyDocumentError(
            "Не удалось прочитать файл. Проверьте формат и структуру.",
        ) from error


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, float):
        if not math.isfinite(value):
            raise SurveyDocumentError("Файл содержит неконечное числовое значение.")
        return str(int(value)) if value.is_integer() else str(value)
    return str(value)


def checked_workbook(sheets: list[SurveySheet], warnings: list[str]) -> SurveyWorkbook:
    if not sheets or len(sheets) > 20:
        raise SurveyDocumentError("В файле должно быть от 1 до 20 листов.")
    cells = 0
    text_size = 0
    for sheet in sheets:
        if len(sheet.rows) > MAX_ROWS:
            raise SurveyDocumentError("Превышен предел 25 000 строк на лист. Данные не обрезаны.")
        for row in sheet.rows:
            if len(row) > MAX_COLUMNS:
                raise SurveyDocumentError("Превышен предел 512 столбцов. Данные не обрезаны.")
            cells += len(row)
            text_size += sum(len(value) for value in row)
            if cells > MAX_CELLS or text_size > MAX_TEXT:
                raise SurveyDocumentError(
                    "Слишком большой набор: до 150 000 ячеек и 2 млн символов.",
                )
    if not any(value.strip() for sheet in sheets for row in sheet.rows for value in row):
        raise SurveyDocumentError("В файле нет заполненных ячеек.")
    return SurveyWorkbook(sheets=sheets, warnings=warnings)


def _parse_workbook(
    filename: str, content: bytes, *, encoding: Literal["utf-8-sig", "cp1251"] = "utf-8-sig",
    delimiter: Literal[",", ";", "\t"] = ",",
) -> SurveyWorkbook:
    if not content or len(content) > MAX_FILE_BYTES:
        raise SurveyDocumentError("Файл пуст или превышает 10 МиБ.")
    suffix = PurePath(filename).suffix.lower()
    if suffix == ".csv":
        try:
            reader = csv.reader(
                io.StringIO(content.decode(encoding)), delimiter=delimiter, strict=True,
            )
            rows: list[list[str]] = []
            cells = text_size = 0
            for row in reader:
                cells += len(row)
                text_size += sum(len(value) for value in row)
                if (len(rows) >= MAX_ROWS or len(row) > MAX_COLUMNS
                        or cells > MAX_CELLS or text_size > MAX_TEXT):
                    raise SurveyDocumentError("CSV превышает допустимый объём. Данные не обрезаны.")
                rows.append(row)
        except (UnicodeError, csv.Error) as error:
            raise SurveyDocumentError(
                "Не удалось прочитать CSV: проверьте кодировку и разделитель.",
            ) from error
        return checked_workbook([SurveySheet(name="CSV", rows=rows)], [])
    if suffix == ".xlsx":
        checked_zip(content)
        check_coordinates(content)
        with ZipFile(io.BytesIO(content)) as archive:
            if "xl/workbook.xml" not in archive.namelist():
                raise SurveyDocumentError("Файл не является книгой XLSX.")
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False,
                                      keep_links=False)
        try:
            sheets = []
            formula_count = 0
            cells = text_size = 0
            if len(book.worksheets) > 20:
                raise SurveyDocumentError("Допустимо до 20 листов.")
            for sheet in book.worksheets:
                # Untrusted dimensions are not used to allocate an unbounded rectangle.
                sheet.reset_dimensions()
                rows = []
                for row in sheet.iter_rows():
                    values = [cell_text(cell.value) for cell in row]
                    formula_count += sum(cell.data_type == "f" for cell in row)
                    cells += len(values)
                    text_size += sum(len(value) for value in values)
                    if (len(rows) >= MAX_ROWS or len(values) > MAX_COLUMNS
                            or cells > MAX_CELLS or text_size > MAX_TEXT):
                        raise SurveyDocumentError(
                            "XLSX превышает допустимый объём. Данные не обрезаны.",
                        )
                    rows.append(values)
                sheets.append(SurveySheet(name=sheet.title, rows=rows))
            warnings = []
            if formula_count:
                warnings.append(f"Формулы ({formula_count}) сохранены как текст и не вычисляются. "
                                "Для анализа значений загрузите файл без формул.")
            return checked_workbook(sheets, warnings)
        finally:
            book.close()
    if suffix == ".xls":
        book = xlrd.open_workbook(file_contents=content, on_demand=True)
        try:
            sheets = []
            if book.nsheets > 20:
                raise SurveyDocumentError("Допустимо до 20 листов.")
            total = 0
            for sheet in book.sheets():
                total += sheet.nrows * sheet.ncols
                if sheet.nrows > MAX_ROWS or sheet.ncols > MAX_COLUMNS or total > MAX_CELLS:
                    raise SurveyDocumentError("XLS превышает допустимый объём. Данные не обрезаны.")
                rows = []
                for index in range(sheet.nrows):
                    values = []
                    for cell in sheet.row(index):
                        value = cell.value
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(value, book.datemode)
                        values.append(cell_text(value))
                    rows.append(values)
                sheets.append(SurveySheet(name=sheet.name, rows=rows))
            return checked_workbook(sheets, [
                "XLS: используются сохранённые значения; формулы не пересчитываются.",
            ])
        finally:
            book.release_resources()
    raise SurveyDocumentError("Поддерживаются XLSX, XLS и CSV. Макросы не поддерживаются.")
