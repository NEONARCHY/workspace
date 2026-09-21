# ruff: noqa: RUF001
"""Prepare or safely send HR employee cards from a verified tenure workbook.

The workbook remains a local source document and is never copied into the repository.
The default mode only validates and prints the number of rows. Sending requires both
``--apply`` and an access token supplied through an environment variable.

Example:

    $env:HR_IMPORT_TOKEN = "<access token>"
    python scripts/import_hr_tenure_workbook.py `
      "C:\\Users\\...\\2026.xlsx" --apply --api-url https://server/api/v1
"""

from __future__ import annotations

import argparse
import json
import os
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from openpyxl import load_workbook


def _date(value: object, *, field: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        match = re.search(r"(\d{2})\.(\d{2})\.(\d{4})", value)
        if match:
            day, month, year = map(int, match.groups())
            return date(year, month, day)
    raise ValueError(f"{field}: ожидалась дата, получено {value!r}")


def read_rows(source: Path, sheet_name: str) -> tuple[date, list[dict[str, object]]]:
    workbook = load_workbook(source, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"В файле нет листа {sheet_name!r}")
    sheet = workbook[sheet_name]
    anchor = _date(sheet.cell(10, 13).value, field="контрольная дата")
    rows: list[dict[str, object]] = []
    for source_row in range(13, sheet.max_row + 1):
        number, full_name = sheet.cell(source_row, 3).value, sheet.cell(source_row, 4).value
        if not isinstance(number, int) or not isinstance(full_name, str) or not full_name.strip():
            continue
        employment_date = _date(sheet.cell(source_row, 10).value, field=f"строка {source_row}")
        years, months, days = (
            sheet.cell(source_row, 12).value,
            sheet.cell(source_row, 13).value,
            sheet.cell(source_row, 14).value,
        )
        if not all(isinstance(value, int) for value in (years, months, days)):
            raise ValueError(
                f"строка {source_row}: стаж должен быть указан годами, месяцами и днями"
            )
        title = sheet.cell(source_row, 8).value
        rows.append(
            {
                "sourceRow": source_row,
                "fullName": full_name.strip(),
                "jobTitle": title.strip() if isinstance(title, str) and title.strip() else None,
                "employmentDate": employment_date.isoformat(),
                "serviceAnchorDate": anchor.isoformat(),
                "serviceYears": years,
                "serviceMonths": months,
                "serviceDays": days,
                "serviceReason": "Импорт из "
                f"{source.name}, лист {sheet_name}, контрольная дата {anchor.isoformat()}",
            }
        )
    if not rows:
        raise ValueError("В листе не найдены строки сотрудников")
    return anchor, rows


def send_import(api_url: str, token: str, payload: dict[str, Any]) -> dict[str, Any]:
    request = Request(
        f"{api_url.rstrip('/')}/hr/profiles/import",
        data=json.dumps(payload).encode(),
        headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"},
        method="POST",
    )
    try:
        with urlopen(request, timeout=30) as response:
            return json.loads(response.read())
    except HTTPError as error:
        raise SystemExit(
            f"Сервер отклонил импорт: HTTP {error.code} {error.read().decode()}"
        ) from error
    except URLError as error:
        raise SystemExit(f"Не удалось подключиться к Workspace: {error.reason}") from error


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--sheet", default="сентябрь")
    parser.add_argument("--api-url", default=os.getenv("YUKSALISH_API_URL"))
    parser.add_argument("--token-env", default="HR_IMPORT_TOKEN")
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    if not args.workbook.is_file():
        raise SystemExit(f"Файл не найден: {args.workbook}")
    anchor, rows = read_rows(args.workbook, args.sheet)
    payload: dict[str, Any] = {
        "sourceLabel": f"{args.workbook.name}:{args.sheet}:{anchor.isoformat()}",
        "rows": rows,
    }
    print(f"Найдено кадровых карточек: {len(rows)}; контрольная дата: {anchor.isoformat()}")
    if not args.apply:
        print("Это безопасный просмотр. Добавьте --apply для отправки в Workspace.")
        return
    if not args.api_url:
        raise SystemExit("Укажите --api-url или переменную YUKSALISH_API_URL")
    token = os.getenv(args.token_env)
    if not token:
        raise SystemExit(f"Укажите access token в переменной окружения {args.token_env}")
    result = send_import(args.api_url, token, payload)
    print(f"Создано: {result['created']}; уже было импортировано: {result['alreadyImported']}")


if __name__ == "__main__":
    try:
        main()
    except ValueError as error:
        raise SystemExit(f"Не удалось прочитать таблицу: {error}") from error
